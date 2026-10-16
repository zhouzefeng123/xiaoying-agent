from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.shared import Inches
from PIL import Image, ImageDraw


def build_image(image_path: Path) -> None:
    image = Image.new("RGB", (320, 120), color="#f4e6bf")
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((8, 8, 312, 112), radius=18, outline="#7a5224", width=4, fill="#fff8e7")
    draw.text((28, 44), "Auramate go-word", fill="#5c3c1c")
    image.save(image_path, format="PNG")


def build_document(output_path: Path) -> None:
    asset_dir = output_path.parent
    asset_dir.mkdir(parents=True, exist_ok=True)
    image_path = asset_dir / "goword-inline-image.png"
    build_image(image_path)

    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("图片前文本")
    picture_run = paragraph.add_run()
    picture_run.add_picture(str(image_path), width=Inches(1.35))
    paragraph.add_run("图片后文本")

    document.add_paragraph("第二段保持普通文本。")
    document.save(output_path)


def main() -> int:
    if len(sys.argv) > 1:
        output_path = Path(sys.argv[1]).expanduser().resolve()
    else:
        output_path = Path("backend/workspace/goword-real-image.docx").resolve()

    build_document(output_path)
    print(output_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())