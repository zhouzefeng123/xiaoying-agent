from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


def build_document(output_path: Path) -> None:
    document = Document()

    title = document.add_heading("Auramate go-word 表格回归文档", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    intro = document.add_paragraph()
    intro_run = intro.add_run("该文档由 python-docx 生成，用于验证 go-word 的真实 docx 打开、表格单元格精确编辑与保存回写。")
    intro_run.font.size = Pt(11)

    table = document.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).paragraphs[0].add_run("原始单元格")
    table.cell(0, 1).paragraphs[0].add_run("第二列")
    table.cell(1, 0).paragraphs[0].add_run("第三格")
    table.cell(1, 1).paragraphs[0].add_run("第四格")

    tail = document.add_paragraph("表格后正文。")
    tail.alignment = WD_ALIGN_PARAGRAPH.LEFT

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def main() -> int:
    if len(sys.argv) > 1:
        output_path = Path(sys.argv[1]).expanduser().resolve()
    else:
        output_path = Path("backend/workspace/goword-real-table.docx").resolve()

    build_document(output_path)
    print(output_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())